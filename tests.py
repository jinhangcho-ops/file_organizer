"""검사: 임시 폴더에 지어낸 파일을 만들어 정리·되돌리기·색인·검색·나무·모아 보기·예외·서버를 확인한다.
실제 다운로드 폴더와 설정 파일은 건드리지 않는다(APPDATA를 임시 폴더로 바꾼다).
다른 디스크로 옮기기는 환경 변수 OTHER_DRIVE(예: D:\\)를 주면 그 디스크에 임시 폴더를 만들어 확인하고 지운다.
실행: python tests.py   실패가 있으면 이름을 보여 주고 종료 코드 1."""
import ctypes
import os
import shutil
import sqlite3
import sys
import tempfile
import time
from pathlib import Path

TEMP = Path(tempfile.mkdtemp(prefix="file_organizer_test_"))
os.environ["APPDATA"] = str(TEMP / "appdata")

import actions  # noqa: E402  APPDATA를 바꾼 뒤에 불러온다.
import exclude  # noqa: E402
import extract  # noqa: E402
import filetime  # noqa: E402
import index  # noqa: E402
import organize  # noqa: E402
import kinds  # noqa: E402
import search  # noqa: E402
import server  # noqa: E402
import settings  # noqa: E402
import tree  # noqa: E402

failures = []
OLD = int(time.mktime((2020, 1, 15, 12, 0, 0, 0, 0, -1)))  # 오래된 시각(2020-01-15 12:00)


def check(name, ok):
    print(("통과 " if ok else "실패 ") + name)
    if not ok:
        failures.append(name)


def raises(name, work, text):
    try:
        work()
    except ValueError as e:
        check(name, text in str(e))
        return
    check(name + " (오류가 나지 않음)", False)


def write(path, text="", encoding="utf-8"):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding=encoding)
    return str(path)


def mark_downloaded(path):
    with open(path + ":Zone.Identifier", "w") as stream:
        stream.write("[ZoneTransfer]\nZoneId=3\n")


def settle():
    """새로 만든 파일은 1초쯤 안에 다른 프로그램(백신 검사로 보인다)이 연 시각을 바꾼다. 시각을 바꾸기 전에 기다린다."""
    time.sleep(1.5)


def make_old(path, accessed=OLD, modified=OLD):
    os.utime(path, (accessed, modified))


def make_pdf(path, text):
    """글자 한 줄이 든 가장 작은 PDF."""
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET"
    parts = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
             "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
             f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, part in enumerate(parts, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{part}\nendobj\n".encode()
    start = len(out)
    out += f"xref\n0 {len(parts) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(parts) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    Path(path).write_bytes(out)
    return str(path)


def make_office(root):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    document = Document()
    document.add_paragraph("계약 조건 검토 의견")
    document.save(root / "계약.docx")
    book = Workbook()
    book.active.append(["매출 합계", 1200])
    book.save(root / "실적.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "발표 자료 요약"
    slides.save(root / "발표.pptx")


def nonzero(pairs):
    return {name: n for name, n in pairs if n}


def names(rows):
    return {row["name"] for row in rows}


def test_organize():
    src, dest = TEMP / "src", TEMP / "dest"
    make_pdf(write(src / "보고서.pdf"), "quarterly zebra report")
    for name in ["사진.JPG", "영상.mp4", "묶음.zip", "setup.exe", "이상한.xyz", "받는중.crdownload", "숨김.txt", "sub/안쪽.txt"]:
        write(src / name, name)
    ctypes.windll.kernel32.SetFileAttributesW(str(src / "숨김.txt"), 2)  # 숨김
    write(dest / "PDF" / "보고서.pdf", "먼저 있던 파일")

    moves = organize.plan(str(src), str(dest))
    to = {Path(m["from"]).name: m["to"] for m in moves}
    check("정리: 옮길 파일 6개(받는 중·숨김·하위 폴더 제외)", sorted(to) == sorted(["보고서.pdf", "사진.JPG", "영상.mp4", "묶음.zip", "setup.exe", "이상한.xyz"]))
    check("정리: 같은 이름은 (2)", to.get("보고서.pdf") == str(dest / "PDF" / "보고서 (2).pdf"))
    check("정리: 종류별 폴더(대문자 확장자 포함)", to.get("사진.JPG") == str(dest / "사진" / "사진.JPG")
          and to.get("setup.exe") == str(dest / "설치파일" / "setup.exe") and to.get("이상한.xyz") == str(dest / "기타" / "이상한.xyz"))
    check("종류: 표·PDF·문서 나누기", kinds.of(".CSV") == "표" and kinds.of(".pdf") == "PDF" and kinds.of(".hwp") == "문서"
          and kinds.of(".JPG") == "사진" and kinds.of(".png") == "이미지")
    check("정리: 미리보기는 옮기지 않음", (src / "영상.mp4").exists())

    result = organize.run(str(src), str(dest))
    check("정리: 6개 옮김", result == {"moved": 6, "failed": []})
    left = sorted(p.name for p in src.iterdir())
    check("정리: 남은 것은 받는 중·숨김·하위 폴더", left == sorted(["받는중.crdownload", "숨김.txt", "sub"]) and (src / "sub" / "안쪽.txt").exists())
    check("정리: 먼저 있던 파일은 그대로", (dest / "PDF" / "보고서.pdf").read_text(encoding="utf-8") == "먼저 있던 파일")

    result = organize.undo()
    check("되돌리기: 6개", result == {"moved": 6, "failed": []} and (src / "영상.mp4").exists() and not (dest / "동영상" / "영상.mp4").exists())
    check("되돌리기: 먼저 있던 파일은 남음", (dest / "PDF" / "보고서.pdf").exists() and (src / "보고서.pdf").exists())
    raises("되돌리기: 기록이 없으면 거부", organize.undo, "되돌릴 정리가 없습니다.")

    organize.run(str(src), str(dest))
    write(src / "영상.mp4", "그사이 새로 생긴 파일")
    result = organize.undo()
    check("되돌리기: 원래 자리에 같은 이름이 생기면 그것만 남김", result["moved"] == 5 and len(result["failed"]) == 1
          and (src / "영상.mp4").read_text(encoding="utf-8") == "그사이 새로 생긴 파일" and (dest / "동영상" / "영상.mp4").exists())
    check("되돌리기: 못 한 것은 기록에 남김", [Path(m["to"]).name for m in organize.last()] == ["영상.mp4"])
    os.remove(src / "영상.mp4")
    check("되돌리기: 다시 하면 남은 것도", organize.undo()["moved"] == 1 and (src / "영상.mp4").exists())

    raises("정리: 없는 폴더", lambda: organize.plan(str(TEMP / "없음"), str(dest)), "정리할 곳: 없는 폴더입니다.")
    raises("정리: 놓을 곳 안 고름", lambda: organize.plan(str(src), ""), "놓을 곳: 폴더를 고르세요.")


def test_other_drive():
    other = os.environ.get("OTHER_DRIVE")
    if not other:
        print("건너뜀 다른 디스크로 옮기기 (OTHER_DRIVE 없음)")
        return
    folder = Path(tempfile.mkdtemp(prefix="file_organizer_test_", dir=other))
    try:
        src = write(TEMP / "drive" / "받은.zip", "압축 내용" * 100)
        mark_downloaded(src)
        settle()
        filetime.set_times(src, OLD * 10**9, OLD * 10**9, OLD * 10**9)
        before = os.stat(src)
        target = str(folder / "받은.zip")
        organize.move(src, target)
        after = os.stat(target)
        check("다른 디스크: 원본 지움", not os.path.exists(src))
        check("다른 디스크: 다운로드 표시 유지", filetime.downloaded(target))
        check("다른 디스크: 만든·고친 시각 유지", after.st_birthtime_ns == before.st_birthtime_ns and after.st_mtime_ns == before.st_mtime_ns)
        check("다른 디스크: 크기 같음", after.st_size == before.st_size)

        again = write(TEMP / "drive" / "받은.zip", "다시")
        done, failed = organize.apply([(again, target)])
        check("다른 디스크: 같은 이름이 있으면 실패하고 원본 둠", not done and failed and os.path.exists(again))

        src_dir = TEMP / "drive2"
        write(src_dir / "문서.txt", "내용")
        result = organize.run(str(src_dir), str(folder))
        check("다른 디스크: 정리", result["moved"] == 1 and (folder / "문서" / "문서.txt").exists())
    finally:
        shutil.rmtree(folder, ignore_errors=True)


def test_index_search():
    root = TEMP / "root"
    memo = write(root / "memo.txt", "회의록: 예산 삼천만원 확정")
    write(root / "옛메모.txt", "한글 메모 cp949", encoding="cp949")
    make_pdf(root / "report.pdf", "quarterly zebra report")
    make_office(root)
    big = root / "big.bin"
    big.write_bytes(b"\0" * (3 * 1024 * 1024))
    make_old(str(big))
    write(root / "sub" / "deep" / "a.txt", "ab cd")
    write(root / "sub2" / "b.txt", "옆 폴더")
    write(root / "AppData" / "skip.txt", "들어가지 않음")
    write(root / "숨긴폴더" / "skip.txt", "들어가지 않음")
    ctypes.windll.kernel32.SetFileAttributesW(str(root / "숨긴폴더"), 2)
    settle()
    make_old(str(big))
    make_old(memo, accessed=OLD, modified=int(time.time()))
    atime = os.stat(memo).st_atime_ns

    settings.save({"roots": [str(root)], "dest": str(TEMP / "dest")})
    paths = index.roots(settings.load())
    check("색인: 폴더 두 개(설정 + 놓을 곳)", sorted(paths) == sorted([str(root), str(TEMP / "dest")]))
    index.build(paths)
    check("색인: 연 시각을 읽기 전으로 되돌림", os.stat(memo).st_atime_ns == atime)

    find = lambda **q: names(search.listing(q)["rows"])  # noqa: E731
    check("검색: 세 글자 이상 내용", find(text="삼천만") == {"memo.txt"})
    check("검색: 두 글자 내용", find(text="예산") == {"memo.txt"})
    check("검색: cp949 내용", find(text="한글 메모") == {"옛메모.txt"})
    check("검색: PDF 내용(대소문자 무시)", find(text="ZEBRA") == {"report.pdf"})
    check("검색: docx·xlsx·pptx 내용", find(text="계약 조건") == {"계약.docx"} and find(text="매출 합계") == {"실적.xlsx"}
          and find(text="발표 자료") == {"발표.pptx"})
    check("검색: 내용 미리보기", "삼천만" in search.listing({"text": "삼천만"})["rows"][0]["snippet"])
    check("검색: 이름(대소문자 무시)", find(name="MEMO") == {"memo.txt"})
    check("검색: 폴더 이름", find(name="deep") == {"a.txt"})
    check("검색: 종류", find(kind="문서", name="발표") == {"발표.pptx"} and find(kind="압축", name="memo") == set())
    listed = search.listing({"name": "memo"})
    check("목록: 개수·합계·종류별 개수", listed["total"] == 1 and nonzero(listed["kinds"]) == {"문서": 1} and [k for k, _ in listed["kinds"]] == kinds.NAMES and listed["bytes"] == os.path.getsize(memo))
    check("목록: 위치 이름", {r["place"] for r in listed["rows"]} == {"root"})
    by_size = [r["name"] for r in search.listing({"folder": str(root), "sort": "size", "order": "desc"})["rows"]]
    check("목록: 크기 정렬", by_size[0] == "big.bin")
    check("목록: 폴더 범위(하위 포함)", find(folder=str(root / "sub")) == {"a.txt"} and find(folder=str(root / "sub2")) == {"b.txt"})
    check("검색: 크기", find(size_min=2) == {"big.bin"} and "big.bin" not in find(size_max=1, name="b"))
    check("검색: 날짜 기간(끝 날짜 포함)", find(date_from="2020-01-01", date_to="2020-01-15") == {"big.bin"} and find(date_to="2020-01-14") == set())
    check("색인: AppData·숨김 폴더 제외", find(name="skip") == set())
    raises("검색: 잘못된 날짜", lambda: search.listing({"date_from": "2020/01/01"}), "날짜는 2026-10-17 모양으로")
    raises("검색: 음수 크기", lambda: search.listing({"size_min": -1}), "최소 크기는 0 이상")
    raises("검색: 모르는 종류", lambda: search.listing({"kind": "그림"}), "모르는 종류입니다: 그림")
    raises("목록: 모르는 목록", lambda: search.listing({"view": "all"}), "모르는 목록입니다")
    check("목록: 조건 없으면 전체", search.listing({})["total"] == index.count())

    write(root / "memo.txt", "회의록: 예산 사천만원으로 변경")
    os.remove(root / "sub2" / "b.txt")
    index.build(paths)
    check("색인 갱신: 바뀐 내용", find(text="사천만") == {"memo.txt"} and find(text="삼천만") == set())
    check("색인 갱신: 지운 파일 빠짐", find(name="b.txt") == set())
    write(root / "sub2" / "b.txt", "옆 폴더")
    index.build(paths)

    top = tree.places()
    check("위치: 색인한 폴더(설정 순서, 놓을 곳은 끝)", [f["path"] for f in top] == paths)
    files_under = sum(1 for p in root.rglob("*") if p.is_file() and not {"AppData", "숨긴폴더"} & set(p.relative_to(root).parts))
    root_row = next(f for f in top if f["path"] == str(root))
    check("위치: 하위 전체 파일 수", root_row["files"] == files_under)
    check("경로: 맨 위 폴더부터", [f["path"] for f in tree.chain(str(root / "sub" / "deep" / "a.txt"))] == [str(root), str(root / "sub"), str(root / "sub" / "deep")])
    check("경로: 색인에 없는 경로", tree.chain(str(TEMP / "없음" / "x.txt")) == [])
    kids = tree.children(str(root))
    check("폴더 나무: 바로 아래 폴더만, 이름 순", [f["name"] for f in kids] == ["sub", "sub2"])
    check("폴더 나무: 하위 폴더 있음 표시", {f["name"]: bool(f["has_children"]) for f in kids} == {"sub": True, "sub2": False}
          and root_row["has_children"])
    check("폴더 나무: 폴더 수 세기(하위 포함)", next(f for f in kids if f["name"] == "sub")["files"] == 1)
    check("경로: 폴더 자신까지", [f["name"] for f in tree.chain(str(root / "sub" / "deep"), folder=True)] == ["root", "sub", "deep"])

    src, dest = TEMP / "src", TEMP / "dest"
    organize.run(str(src), str(dest))
    check("정리 후 색인: 옮긴 파일이 그 폴더 목록에", "보고서 (2).pdf" in find(folder=str(dest / "PDF")))
    check("정리 후 색인: 내용으로 찾음", str(dest / "PDF" / "보고서 (2).pdf") in {r["path"] for r in search.listing({"text": "zebra"})["rows"]})
    dest_row = next(f for f in tree.places() if f["path"] == str(dest))
    check("정리 후 색인: 폴더 합계", dest_row["files"] == sum(1 for p in dest.rglob("*") if p.is_file()))
    organize.undo()
    check("되돌린 후 색인: 색인 밖으로 간 파일은 빠짐", str(dest / "PDF" / "보고서 (2).pdf") not in {r["path"] for r in search.listing({"text": "zebra"})["rows"]})


def test_recent():
    root = TEMP / "root"
    got = write(root / "받은파일.zip", "zip")
    mark_downloaded(got)
    old = write(root / "옛날.txt", "옛날")
    settle()
    make_old(old)
    index.build(index.roots(settings.load()))

    rows = lambda view: names(search.listing({"view": view})["rows"])  # noqa: E731
    check("모아 보기: 받은 파일", rows("downloaded") == {"받은파일.zip"})
    check("모아 보기: 고친 파일", "memo.txt" in rows("modified") and "옛날.txt" not in rows("modified"))
    check("모아 보기: 오래 안 연 파일", "옛날.txt" in rows("stale") and "big.bin" in rows("stale") and "memo.txt" not in rows("stale"))
    tags = {r["name"]: r["stale"] for r in search.listing({"folder": str(root)})["rows"]}
    check("목록: 오래 안 연 태그", tags["옛날.txt"] and not tags["memo.txt"])
    settings.save({"stale_days": 3650})
    check("모아 보기: 기간 설정을 따름", "옛날.txt" not in rows("stale"))
    settings.save({"stale_days": 90})
    counts = search.counts()
    check("사이드바 숫자", counts["downloaded"] == 1 and counts["all"] == index.count() and counts["stale"] == len(rows("stale")))

    exclude.add(str(root / "sub"))
    exclude.add(old)
    check("예외: 폴더 아래 전부 빠짐", "a.txt" not in rows("modified"))
    check("예외: 이름이 이어지는 옆 폴더는 남음", "b.txt" in rows("modified"))
    check("예외: 파일 하나", "옛날.txt" not in rows("stale"))
    check("예외: 검색에는 나옴", {"a.txt", "옛날.txt"} <= names(search.listing({"name": "."})["rows"]))
    exclude.remove(str(root / "sub"))
    exclude.remove(old)
    check("예외 빼기: 다시 나옴", "a.txt" in rows("modified") and "옛날.txt" in rows("stale"))
    raises("예외: 없는 경로", lambda: exclude.add(str(TEMP / "없음")), "없는 파일·폴더입니다")


def test_duplicates():
    root = TEMP / "root" / "dup"
    one = write(root / "같은1.txt", "같은 내용입니다")
    write(root / "같은2.txt", "같은 내용입니다")
    write(root / "크기만같음.txt", "다른 내용입니다")
    big = b"x" * (9 * 1024 * 1024)
    (root / "큰1.bin").write_bytes(big)
    (root / "큰2.bin").write_bytes(big)
    settle()
    make_old(one, accessed=OLD, modified=int(time.time()))
    atime = os.stat(one).st_atime_ns
    index.build(index.roots(settings.load()))
    found = search.listing({"view": "duplicates", "folder": str(root)})
    check("중복: 내용이 같은 것만", names(found["rows"]) == {"같은1.txt", "같은2.txt", "큰1.bin", "큰2.bin"})
    check("중복: 태그", all(r["duplicate"] for r in found["rows"]) and not any(r["duplicate"] for r in search.listing({"name": "크기만같음"})["rows"]))
    check("중복: 지문을 만들어도 연 시각은 그대로", os.stat(one).st_atime_ns == atime)
    views = dict(search.listing({"folder": str(root)})["views"])
    check("필터 칩 숫자: 지금 위치 안에서", views["duplicates"] == 4 and set(views) == set(search.VIEWS))
    views = dict(search.listing({"folder": str(root), "kind": "기타"})["views"])
    check("필터 칩 숫자: 고른 종류 안에서", views["duplicates"] == 2)
    both = search.listing({"folder": str(root), "view": "duplicates", "kind": "문서"})
    check("필터 칩 + 종류 칩 함께", names(both["rows"]) == {"같은1.txt", "같은2.txt"} and nonzero(both["kinds"]) == {"문서": 2, "기타": 2})
    write(root / "같은2.txt", "이제 다른 내용이에요")
    index.build(index.roots(settings.load()))
    check("중복: 바뀌면 다시 비교", names(search.listing({"view": "duplicates", "folder": str(root)})["rows"]) == {"큰1.bin", "큰2.bin"})


def test_stop():
    paths = index.roots(settings.load())
    gone = TEMP / "root" / "sub2" / "b.txt"
    os.remove(gone)
    places = tree.places()
    index.stop.set()
    try:
        index.build(paths)
    finally:
        index.stop.clear()
    check("멈춤: 끝까지 못 훑었으면 사라진 파일을 지우지 않음", "b.txt" in names(search.listing({"folder": str(gone.parent)})["rows"]))
    check("멈춤: 위치(폴더 목록)는 그대로", tree.places() == places)
    index.build(paths)
    check("멈춤 뒤 다시 하면 정리됨", "b.txt" not in names(search.listing({"folder": str(gone.parent)})["rows"]))
    write(gone, "옆 폴더")

    real = index.walk

    def walk_then_stop(root):  # 첫 폴더를 넘겨준 뒤 [멈추기]를 누른 것처럼 한다.
        for n, item in enumerate(real(root)):
            if n == 1:
                client.post("/api/index/stop", json={})  # 열쇠 쿠키를 받은 클라이언트
            yield item

    client = server.app.test_client()
    client.get("/?key=" + server.KEY)
    index.walk = walk_then_stop
    try:
        client.post("/api/index", json={})
        for _ in range(100):
            if not index.status["running"]:
                break
            time.sleep(0.1)
    finally:
        index.walk = real
    check("멈춤: 화면의 [멈추기]", index.status["stopped"] and not index.status["running"] and not index.status["error"])
    index.stop.clear()  # 서버로 보낸 멈춤 신호를 지운다(다음 색인은 index.start가 지우지만 여기서는 build를 바로 부른다).
    index.build(paths)


def test_resume():
    root = TEMP / "resume"
    for n in range(6):
        write(root / f"메모{n}.txt", f"이어 읽기 검사 {n}번")
    settings.save({"roots": settings.load()["roots"] + [str(root)]})
    paths = index.roots(settings.load())
    index.build(paths)  # 지금 있는 것을 모두 읽어 둔다.

    reads = []
    real = extract.text

    def counting(path, ext):  # 내용을 읽을 때마다 센다.
        reads.append(os.path.basename(path))
        return real(path, ext)

    extract.text = counting
    try:
        index.build(paths)
        check("다시 색인: 한 번 읽은 파일은 다시 읽지 않음", reads == [])

        write(root / "메모0.txt", "바뀐 내용입니다")
        for n in range(6, 12):
            write(root / f"메모{n}.txt", f"이어 읽기 검사 {n}번")
        reads.clear()

        def stop_after_three(path, ext):  # 세 개를 읽은 뒤 [멈추기]를 누른 것처럼 한다.
            if len(reads) == 3:
                index.stop.set()
            return counting(path, ext)

        extract.text = stop_after_three
        try:
            index.build(paths)
        finally:
            index.stop.clear()
        first = list(reads)
        db = sqlite3.connect(settings.folder() / "index.db")  # 다른 연결로 봐도 저장돼 있어야 한다.
        stored = db.execute("select count(*) from files where textstamp = stamp and path like ?", (str(root) + "%",)).fetchone()[0]
        db.close()
        check("멈춤: 읽은 데까지 저장됨", len(first) in (3, 4) and stored == 12 - 7 + len(first))
        extract.text = counting
        reads.clear()
        index.build(paths)
        check("이어서: 남은 것만 읽음", len(reads) == 7 - len(first) and not set(reads) & set(first))
        check("이어서: 바뀐 파일은 다시 읽음", "메모0.txt" in first + reads)
        check("이어서: 내용 검색", names(search.listing({"text": "바뀐 내용"})["rows"]) == {"메모0.txt"}
              and names(search.listing({"text": "검사 11번"})["rows"]) == {"메모11.txt"})
    finally:
        extract.text = real


def find_in(folder):
    return names(search.listing({"folder": str(folder)})["rows"])


def test_actions():
    root = TEMP / "root" / "act"
    first = write(root / "첫째.txt", "하나")
    write(root / "다른.txt", "둘")
    write(root / "옮길곳" / "다른.txt", "같은 이름")
    index.build(index.roots(settings.load()))

    renamed = actions.rename(first, "바꾼이름.md")
    check("이름 바꾸기", renamed == str(root / "바꾼이름.md") and os.path.exists(renamed) and not os.path.exists(first))
    row = search.listing({"name": "바꾼이름"})["rows"]
    check("이름 바꾸기: 색인(이름·종류)", names(row) == {"바꾼이름.md"} and row[0]["kind"] == "문서" and not search.listing({"name": "첫째"})["rows"])
    raises("이름 바꾸기: 같은 이름", lambda: actions.rename(renamed, "다른.txt"), "같은 이름의 파일이 있습니다")
    raises("이름 바꾸기: 쓸 수 없는 글자", lambda: actions.rename(renamed, "a/b.txt"), "쓸 수 없는 이름입니다")
    raises("이름 바꾸기: 빈 이름", lambda: actions.rename(renamed, "  "), "쓸 수 없는 이름입니다")
    renamed = actions.rename(renamed, "바꾼이름.MD")
    check("이름 바꾸기: 대소문자만", os.path.basename(renamed) == "바꾼이름.MD" and "바꾼이름.MD" in os.listdir(root))

    moved = actions.move_to(renamed, str(root / "옮길곳"))
    check("옮기기", moved == str(root / "옮길곳" / "바꾼이름.MD") and os.path.exists(moved) and not os.path.exists(renamed))
    check("옮기기: 색인", find_in(root / "옮길곳") == {"바꾼이름.MD", "다른.txt"})
    raises("옮기기: 같은 이름", lambda: actions.move_to(str(root / "다른.txt"), str(root / "옮길곳")), "같은 이름의 파일이 있습니다")
    raises("옮기기: 이미 그 폴더", lambda: actions.move_to(moved, str(root / "옮길곳")), "이미 그 폴더에 있습니다")
    raises("옮기기: 없는 폴더", lambda: actions.move_to(moved, str(root / "없음")), "없는 폴더입니다")

    calls = []

    def fake_delete(operation):  # 실제 휴지통 대신: 넘어온 경로와 설정을 기록하고 파일을 지운다.
        op = ctypes.cast(operation, ctypes.POINTER(actions.FileOp)).contents
        calls.append((op.pFrom, op.fFlags))
        os.remove(op.pFrom)
        return 0

    real, actions.shell_delete = actions.shell_delete, fake_delete
    try:
        actions.trash(moved)
    finally:
        actions.shell_delete = real
    check("휴지통: 되살릴 수 있게 보냄(영구 삭제면 경고)", bool(calls) and calls[0][0] == moved
          and calls[0][1] & actions.FOF_ALLOWUNDO and calls[0][1] & actions.FOF_WANTNUKEWARNING)
    check("휴지통: 색인에서 빠짐", find_in(root / "옮길곳") == {"다른.txt"})


def test_server():
    client = server.app.test_client()
    check("서버: 열쇠 없으면 403", client.get("/api/state").status_code == 403 and client.get("/").status_code == 403
          and client.get("/static/app.js").status_code == 403)
    page = client.get("/?key=" + server.KEY)
    check("서버: 열쇠로 화면", page.status_code == 200 and "script-src 'self'" in page.headers["Content-Security-Policy"])
    state = client.get("/api/state")
    check("서버: 상태", state.status_code == 200 and state.json["counts"]["all"] > 0 and "기타" in state.json["kinds"] and bool(state.json["places"]))
    found = client.post("/api/files", json={"text": "zebra"})
    check("서버: 목록", found.status_code == 200 and "report.pdf" in names(found.json["rows"]))
    kids = client.get("/api/tree?path=" + str(TEMP / "root"))
    check("서버: 폴더 나무", kids.status_code == 200 and "sub" in {f["name"] for f in kids.json["folders"]})
    bad = client.post("/api/rename", json={"path": str(TEMP / "없음.txt"), "name": "x.txt"})
    check("서버: 색인에 없는 경로는 이름 바꾸지 않음", bad.status_code == 400 and "색인에 없는 경로" in bad.json["error"])
    bad = client.post("/api/trash", json={"path": str(TEMP / "src")})
    check("서버: 색인에 없는 경로는 휴지통에 보내지 않음", bad.status_code == 400 and os.path.exists(TEMP / "src"))
    bad = client.post("/api/move", json={"path": str(TEMP / "root" / "memo.txt")})
    check("서버: 창 없이 옮기기 거부", bad.status_code == 400 and os.path.exists(TEMP / "root" / "memo.txt"))
    bad = client.post("/api/files", json={"view": "all"})
    check("서버: 모르는 목록", bad.status_code == 400)
    bad = client.post("/api/open", json={"path": str(TEMP / "없음.txt")})
    check("서버: 색인에 없는 경로는 열지 않음", bad.status_code == 400 and "색인에 없는 경로" in bad.json["error"])
    bad = client.post("/api/settings", json={"recent_days": 0})
    check("서버: 잘못된 기간", bad.status_code == 400 and "1~3650일" in bad.json["error"])
    bad = client.post("/api/settings", json={"color": "red"})
    check("서버: 모르는 설정", bad.status_code == 400 and "모르는 설정입니다: color" in bad.json["error"])
    bad = client.post("/api/settings", json={"stale_days": "많이"})
    check("서버: 기간이 숫자가 아님", bad.status_code == 400)
    ok = client.post("/api/settings", json={"stale_days": 30})
    check("서버: 기간 저장", ok.status_code == 200 and settings.load()["stale_days"] == 30)
    bad = client.post("/api/pick", json={})
    check("서버: 창 없이 폴더 고르기 거부", bad.status_code == 400)
    bad = client.post("/api/organize/preview", json={"source": str(TEMP / "없음"), "dest": str(TEMP / "dest")})
    check("서버: 없는 정리할 곳", bad.status_code == 400 and "정리할 곳: 없는 폴더입니다." in bad.json["error"])
    started = client.post("/api/index", json={})
    check("서버: 색인 시작", started.status_code == 200 and started.json["running"])
    for _ in range(100):
        if not index.status["running"]:
            break
        time.sleep(0.1)
    check("서버: 색인 끝", not index.status["running"] and not index.status["error"])


try:
    test_organize()
    test_other_drive()
    test_index_search()
    test_recent()
    test_duplicates()
    test_actions()
    test_stop()
    test_resume()
    test_server()
finally:
    shutil.rmtree(TEMP, ignore_errors=True)

print(f"\n실패 {len(failures)}개: {', '.join(failures)}" if failures else "\n모두 통과")
sys.exit(1 if failures else 0)
