"""바이트가 달라도 같은 파일 찾기: 문서 정보만 다른 Office 문서(doc), 글자가 같은 문서(text), 보이는 내용이 같은 사진·동영상(media).
그 파일과 같은 것만 보기(same_as), 색인 단계 순서·끄기·한 단계만 돌리기, 무거운 라이브러리를 켤 때 올리지 않기.
문서·사진·동영상은 모두 테스트 안에서 지어낸다(동영상은 ffmpeg의 시험 화면)."""
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image, ImageDraw

import index
import media
import settings
from conftest import build

TEXT = "올해 사업 계획: 첫째 분기에는 새 지점을 열고, 둘째 분기에는 온라인 판매를 늘린다."


def docx(path, text, bold=False):
    document = Document()
    document.add_paragraph().add_run(text).bold = bold
    document.save(path)
    return str(path)


def change_author(source, target, author):
    """문서 정보(docProps/core.xml)의 만든 사람만 바꾼 사본."""
    with zipfile.ZipFile(source) as old, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as new:
        for item in old.infolist():
            data = old.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = data.replace(b"<dc:creator>", b"<dc:creator>" + author.encode())
            new.writestr(item, data)
    return str(target)


def picture(path, size=(400, 300), seed=0, quality=None):
    img = Image.new("RGB", size, (30, 60, 90))
    draw = ImageDraw.Draw(img)
    w, h = size
    for n in range(6):
        x, y = (37 * (n + 1) * (seed + 3)) % w, (53 * (n + 2) * (seed + 5)) % h
        draw.ellipse((x, y, x + w // 4, y + h // 4), fill=((n * 40 + seed * 70) % 256, 200 - n * 20, (n * 90) % 256))
    img.save(path, **({"quality": quality} if quality else {}))
    return str(path)


def make_video(path, source, size, extra=()):
    subprocess.run([media.ffmpeg(), "-v", "error", "-y", "-f", "lavfi", "-i", f"{source}=duration=40:size={size}:rate=5", *extra, str(path)],
                   check=True, creationflags=media.NO_WINDOW)
    return str(path)


def same_as(client, path):
    rows = client.post("/api/files", json={"same_as": path}).json["rows"]
    return [(row["name"], row["why"]) for row in rows]


def test_office_documents_with_only_document_info_changed_or_same_text(home, client):
    a = docx(home / "계획.docx", TEXT)
    b = change_author(a, home / "계획_사본.docx", "다른 사람")
    docx(home / "계획_굵게.docx", TEXT, bold=True)
    docx(home / "다른 문서.docx", "전혀 다른 내용의 문서입니다. 회의록과 출장 보고를 담았습니다.")
    assert Path(a).read_bytes() != Path(b).read_bytes()
    build(home)
    found = same_as(client, a)  # 고른 파일이 맨 위, 나머지는 정렬 순서(같은 시각이면 id 순)
    assert found[0] == ("계획.docx", "self") and sorted(found[1:]) == [("계획_굵게.docx", "text"), ("계획_사본.docx", "doc")]
    detail = client.get("/api/same?path=" + a).json["content"]
    assert [(f["name"], f["why"]) for f in detail] == [("계획_굵게.docx", "text"), ("계획_사본.docx", "doc")]
    rows = {row["name"]: row["duplicate"] for row in client.post("/api/files", json={"view": "duplicates"}).json["rows"]}
    assert rows == {"계획.docx": "doc", "계획_사본.docx": "doc", "계획_굵게.docx": "text"}  # 센 까닭을 붙인다
    assert client.post("/api/files", json={"same_as": str(home / "없음.docx")}).status_code == 400


def test_short_text_is_not_matched(home, client):
    a = docx(home / "짧음1.docx", "확인")
    docx(home / "짧음2.docx", "확인", bold=True)
    build(home)
    assert same_as(client, a) == [("짧음1.docx", "self")]


def test_steps_can_be_turned_off_and_run_alone(home, client):
    a = docx(home / "계획.docx", TEXT)
    change_author(a, home / "사본.docx", "누군가")
    steps = [{"name": name, "on": name not in ("doc", "media", "recognize")} for name in settings.STEPS]
    assert client.post("/api/settings", json={"index_steps": steps}).status_code == 200
    build(home)
    db = index.connect()
    assert db.execute("select count(*) from files where dochash is not null").fetchone()[0] == 0  # 꺼 둔 단계는 하지 않는다
    db.close()
    assert client.post("/api/index", json={"only": "doc"}).status_code == 200
    index.worker.join(30)
    assert same_as(client, a) == [("계획.docx", "self"), ("사본.docx", "doc")]
    assert client.post("/api/index", json={"only": "모름"}).status_code == 400
    bad = [{"name": "hash", "on": True}]
    assert client.post("/api/settings", json={"index_steps": bad}).status_code == 400
    assert client.post("/api/settings", json={"video_similar": 120}).status_code == 400


def test_resized_picture_is_the_same_but_another_is_not(home, client):
    original = picture(home / "원본.png", (800, 600), seed=1)
    with Image.open(original) as img:
        img.resize((200, 150)).save(home / "작게.jpg", quality=70)
    picture(home / "다른 그림.png", (800, 600), seed=7)
    build(home)
    assert same_as(client, original) == [("원본.png", "self"), ("작게.jpg", "media")]
    assert client.get("/api/same?path=" + original).json["content"][0]["score"] >= 90


def test_reencoded_video_is_the_same_but_another_is_not(home, client):
    settings.save({"video_interval": 30})
    a = make_video(home / "영상.mp4", "testsrc", "320x240")
    make_video(home / "영상_작게.mp4", "testsrc", "160x120", ["-crf", "35"])
    make_video(home / "다른 영상.mp4", "testsrc2", "320x240")
    build(home)
    found = same_as(client, a)
    assert found == [("영상.mp4", "self"), ("영상_작게.mp4", "media")]


def test_video_similarity_rules():
    h = [0x0123456789ABCDEF, 0x1111222233334444, 0x5555666677778888, 0x9999AAAABBBBCCCC, 0xDDDDEEEEFFFF0000]
    assert media.similarity(h, list(h)) == 100
    assert media.similarity(h, [0] + h[:-1]) == 80  # 한 장 밀린 영상: 끝 장면 하나만 짝이 없다(기준 70% 넘음)
    assert media.similarity(h, [None, None, None, None, h[4]]) is None  # 한 가지 색 장면을 빼면 비교할 장면이 3장보다 적다
    assert media.similarity(h, [x ^ 0xFFFF_FFFF for x in h]) == 0
    assert media.dhash(bytes([128] * 72)) is None  # 한 가지 색 장면
    assert media.frame_times(40, 120) == [4.0, 12.0, 20.0, 28.0, 36.0]
    assert media.frame_times(1200, 120)[:2] == [60.0, 180.0] and len(media.frame_times(1200, 120)) == 10


def test_heavy_libraries_are_not_loaded_at_start():
    code = "import sys, server; print(sorted(m for m in ('PIL', 'onnxruntime', 'cv2', 'numpy', 'imageio_ffmpeg', 'tokenizers') if m in sys.modules))"
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=Path(__file__).resolve().parent.parent,
                         env={**os.environ}).stdout.strip()
    assert out == "[]"


@pytest.fixture(autouse=True)
def need_ffmpeg():
    if not shutil.which(media.ffmpeg()):
        pytest.skip("ffmpeg가 없다")


def test_scenes_grabbed_together_match_one_at_a_time(tmp_path, monkeypatch):
    """ffmpeg를 동시에 띄워 뽑은 장면이 하나씩 뽑을 때와 바이트까지 같고, 순서도 시각 순서 그대로다(뽑지 못한 시각은 None)."""
    path = make_video(tmp_path / "a.mp4", "testsrc2", "320x240")
    seconds = [1.5, 37.0, 4.25, 20.0, 999.0]  # 순서를 섞고, 영상 밖의 시각도 넣는다

    def alone(second):
        result = subprocess.run([media.ffmpeg(), "-v", "error", "-ss", f"{second:.2f}", "-i", path, "-frames:v", "1",
                                 "-vf", "scale=9:8,format=gray", "-f", "rawvideo", "-"], capture_output=True, creationflags=media.NO_WINDOW)
        return media.dhash(result.stdout) if result.returncode == 0 else None

    expected = [alone(s) for s in seconds]
    assert expected[4] is None and all(h is not None for h in expected[:4])
    assert media.frame_hashes(path, seconds) == expected
    monkeypatch.setattr(media, "WORKERS", 1)
    assert media.frame_hashes(path, seconds) == expected


def test_matching_resumes_after_an_interrupted_run(home, client, monkeypatch):
    """지문을 만든 뒤 짝을 맞추기 전에 끊겨도(창을 닫는 등) 다음 색인이 이어서 짝을 맞춘다."""
    build(home)  # 한 번 끝까지: 동영상 짝을 맞췄다는 표시(meta)가 남는다
    original = picture(home / "원본.png", (800, 600), seed=1)
    with Image.open(original) as img:
        img.resize((200, 150)).save(home / "작게.jpg", quality=70)
    make_video(home / "영상.mp4", "testsrc2", "320x240")
    real_images, real_videos, calls = media.match_images, media.match_videos, []

    def cut(db, *args):
        db.commit()  # 2초마다 저장(save_now)이 지문을 이미 남긴 뒤에 끊긴 것처럼
        raise RuntimeError("끊김")

    monkeypatch.setattr(media, "match_images", cut)
    with pytest.raises(RuntimeError):
        build(home)
    monkeypatch.setattr(media, "match_images", real_images)
    monkeypatch.setattr(media, "match_videos", lambda db, similar: (calls.append(similar), real_videos(db, similar)))
    build(home)
    assert same_as(client, original) == [("원본.png", "self"), ("작게.jpg", "media")]
    assert calls  # 새 동영상의 짝도 다시 맞췄다
