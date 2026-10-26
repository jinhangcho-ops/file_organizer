"""일곱 번째 감사(색인·검색 쪽)에서 막은 것: 느린 파일을 읽는 동안 잡고 있던 쓰기 잠금, 잠긴 PDF·Office 문서의 글자,
클라우드 전용 파일, 큰 폴더·중복 의심·크기 조건의 조회 방식, 새 동영상만 견주기, 깊은 곳의 게임 이름 폴더, 큰 문서의 글자 읽기,
깨진 주소의 포트, 끝에 점이 붙은 이름, 이름 검색의 범위, 잘못된 단계 목록, 늘 읽을 수 없는 사진.
파일은 모두 지어낸다. 실제 휴지통·클라우드 파일은 쓰지 않는다(클라우드 전용은 판별 함수를 바꿔 끼운다)."""
import ctypes
import os
import random
import sqlite3
import subprocess
import threading
import types

import pytest

import blacklist
import captures
import extract
import games
import index
import lock
import media
import models
import organize
import provenance
import recognize
import search
import settings
import tree
from conftest import build, write


def picture(path, color=(200, 10, 10)):
    from PIL import Image
    img = Image.new("RGB", (40, 30), (10, 200, 10))
    img.paste(color, (0, 0, 20, 30))
    img.save(path)
    return str(path)


def docx(path, paragraphs=("문서 글자입니다",)):
    from docx import Document
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return str(path)


def share_none(path):
    """다른 프로그램이 공유 없이 연 것처럼(열기부터 막힌다). 돌려준 핸들은 닫아야 한다."""
    return ctypes.windll.kernel32.CreateFileW(str(path), 0x80000000, 0, None, 3, 0x80, None)


def names(rows):
    return {os.path.basename(row["path"]) for row in rows}


# B1 ─ 단계마다: 한 파일을 쓴 뒤 다음 파일을 읽는 동안 다른 연결(옮기기·휴지통의 색인 고치기)이 쓸 수 있어야 한다.
class Gate:
    """둘째로 불릴 때 reading을 켜고 go를 기다린 뒤 진짜 함수를 부른다(느린 파일을 읽는 동안처럼)."""

    def __init__(self, real):
        self.real, self.calls, self.reading, self.go = real, 0, threading.Event(), threading.Event()

    def __call__(self, *args, **kwargs):
        self.calls += 1
        if self.calls == 2:
            self.reading.set()
            self.go.wait(30)
        return self.real(*args, **kwargs)


def ready_text(home, patch):
    write(home / "a.txt", "첫째 글자")
    write(home / "b.txt", "둘째 글자")
    build(home)
    gate = Gate(extract.text)
    patch.setattr(extract, "text", gate)
    return gate, index.read_texts, "update files set textstamp = null"


def ready_doc(home, patch):
    docx(home / "a.docx")
    docx(home / "b.docx", ["다른 글자"])
    build(home)
    gate = Gate(index.doc_fingerprint)
    patch.setattr(index, "doc_fingerprint", gate)
    return gate, index.documents, "update files set dochashed = null"


def ready_hash(home, patch):
    write(home / "a.bin", "같은크기1")
    write(home / "b.bin", "같은크기2")
    build(home)
    gate = Gate(index.fingerprint)
    patch.setattr(index, "fingerprint", gate)
    return gate, index.fingerprints, "update files set hashed = null"


def ready_media(home, patch):
    picture(home / "a.png")
    picture(home / "b.png", (10, 10, 200))
    build(home)
    gate = Gate(media.image_hash)
    patch.setattr(media, "image_hash", gate)
    return gate, media.run, "delete from media"


def ready_recognize(home, patch):
    import numpy
    picture(home / "a.png")
    picture(home / "b.png", (10, 10, 200))
    build(home)
    engine = types.SimpleNamespace(np=numpy, faces=None, custom=None, kinds=lambda vector: [], image_vector=lambda image: None, close=lambda: None)
    patch.setattr(models, "ready", lambda name: True)
    patch.setattr(recognize, "Engine", lambda config: engine)
    gate = Gate(recognize.open_image)
    patch.setattr(recognize, "open_image", gate)
    return gate, lambda db: recognize.work(db, lambda **changes: None), "delete from recognized"


@pytest.mark.parametrize("ready", [ready_text, ready_doc, ready_hash, ready_media, ready_recognize])
def test_steps_let_others_write_while_reading_the_next_file(home, monkeypatch, ready):
    with monkeypatch.context() as patch:
        patch.setattr(index, "SAVE_EVERY", 10 ** 6)  # 시간으로는 저장하지 않게(늦은 PC에서도 같은 결과)
        patch.setattr(index, "RELEASE_AT", 0)  # 지어낸 파일은 작으니 큰 파일처럼(읽기 전에 저장)
        gate, step, reset = ready(home, patch)
        db = index.connect()
        db.execute(reset)
        db.commit()
        db.close()
        errors = []

        def work():
            db = index.connect()
            try:
                step(db)
            except Exception as e:  # 검사가 본다
                errors.append(e)
            finally:
                db.close()

        thread = threading.Thread(target=work)
        thread.start()
        try:
            assert gate.reading.wait(30), "둘째 파일을 읽지 않았습니다"
            other = sqlite3.connect(settings.folder() / "index.db", timeout=0.5)
            try:
                other.execute("update files set accessed = accessed")  # 옮기기·휴지통의 index.moved/removed처럼 쓴다
                other.commit()
            finally:
                other.close()
        finally:
            gate.go.set()
            thread.join(60)
        assert not errors, errors


# B2 ─ 잠긴 PDF·Office 문서는 '글자 없음'으로 굳히지 않고 다음에 다시 읽는다. 늘 열 수 없는 문서는 빈 글자로 둔다.
def test_locked_pdf_and_office_documents_are_read_later(tmp_path, unlock_after):
    import pypdfium2
    pdf = tmp_path / "잠긴.pdf"
    document = pypdfium2.PdfDocument.new()
    document.new_page(100, 100)
    document.save(str(pdf))
    document.close()
    word = docx(tmp_path / "잠긴.docx")
    for path, ext in ((str(pdf), ".pdf"), (word, ".docx")):
        handle = share_none(path)
        try:
            assert extract.text(path, ext) is None, ext  # 다음 색인에 다시
        finally:
            ctypes.windll.kernel32.CloseHandle(handle)
        assert extract.text(path, ext) is not None
    sid = lock.user_sid()
    subprocess.run(["icacls", word, "/deny", f"*{sid}:(R)"], check=True, capture_output=True)  # 권한이 아예 없다
    try:
        assert extract.text(word, ".docx") == ""  # 다시 해도 같으니 읽은 것으로
    finally:
        subprocess.run(["icacls", word, "/remove:d", f"*{sid}"], capture_output=True)


# B3 ─ 클라우드 전용 파일: 이름·크기·날짜만 색인하고 내용을 읽는 단계는 모두 건너뛴다(기록도 남기지 않아 받은 뒤에 한다).
def test_cloud_only_attributes():
    assert index.cloud_only(types.SimpleNamespace(st_file_attributes=0x400000 | 0x20))
    assert index.cloud_only(types.SimpleNamespace(st_file_attributes=0x1000))
    assert not index.cloud_only(types.SimpleNamespace(st_file_attributes=0x20 | 0x80000))  # 0x80000: 이 PC에 둠(고정)


def test_cloud_only_files_are_listed_but_not_read(home, client, monkeypatch):
    write(home / "a.txt", "같은크기 글자")
    write(home / "b.txt", "같은크기 글짜")
    docx(home / "c.docx")
    picture(home / "d.png")

    def no_read(*args):
        raise AssertionError(f"클라우드 전용 파일을 읽었습니다: {args}")

    with monkeypatch.context() as patch:
        patch.setattr(index, "cloud_only", lambda st: True)
        for module, name in ((extract, "text"), (index, "doc_fingerprint"), (media, "image_hash"), (media, "video_length")):
            patch.setattr(module, name, no_read)
        build(home)
        assert names(client.post("/api/files", json={}).json["rows"]) == {"a.txt", "b.txt", "c.docx", "d.png"}
        db = index.connect()
        try:
            assert db.execute("select count(*) from files where textstamp is not null or hashed is not null or dochashed is not null").fetchone()[0] == 0
            assert db.execute("select count(*) from media").fetchone()[0] == 0
            row = dict(db.execute("select id, path, ext, stamp from files where name = 'd.png'").fetchone())
            assert recognize.recognize_one(db, types.SimpleNamespace(np=None), row, {}, None) is False
            assert db.execute("select count(*) from recognized").fetchone()[0] == 0
        finally:
            db.close()
    build(home)  # 받은 뒤: 이제 읽는다
    db = index.connect()
    try:
        assert db.execute("select count(*) from files where ext <> '.png' and textstamp is null").fetchone()[0] == 0
        assert db.execute("select count(*) from media").fetchone()[0] == 1
    finally:
        db.close()


# B4 ─ 큰 폴더(전체의 많은 부분)는 key 색인을 한 줄씩 따라가지 않고 훑는다. 작은 폴더는 색인을 탄다.
def plan_of(sql, values):
    db = index.connect()
    try:
        return " ".join(str(r[3]) for r in db.execute("explain query plan " + sql, values))
    finally:
        db.close()


def test_big_folder_is_scanned_small_folder_uses_key_index(home):
    for n in range(9):
        write(home / "큰" / f"{n}.txt")
    write(home / "작은" / "a.txt")
    build(home)
    for folder, uses_index in ((home / "큰", False), (home / "작은", True), (home, False)):
        part, values = search.under(str(folder))
        plan = plan_of("select f.path from files f where " + part, values)
        assert ("sqlite_autoindex_files_1" in plan) == uses_index, (folder, plan)
        assert names(search.listing({"folder": str(folder)})["rows"]) == ({"a.txt"} if uses_index else {f"{n}.txt" for n in range(9)} | ({"a.txt"} if folder == home else set()))
    assert tree.summary(str(home / "큰"))["kinds"] == [["문서", 9]]


# B5 ─ 중복 의심(지문 묶기 세 가지)은 목록 한 번에 한 번만 센다(목록·종류별 개수·합계마다 다시 세지 않게).
def test_duplicate_groups_are_counted_once_per_listing(home, monkeypatch):
    write(home / "a.txt", "같은 내용")
    write(home / "b" / "a.txt", "같은 내용")
    write(home / "c.txt", "다른 내용!")
    build(home)
    statements = []
    real = index.connect

    def traced():
        db = real()
        db.set_trace_callback(statements.append)
        return db

    with monkeypatch.context() as patch:
        patch.setattr(index, "connect", traced)
        found = search.listing({"view": "duplicates", "kind": "문서"})
        summary = tree.summary(str(home))
        counts = search.counts()
    assert names(found["rows"]) == {"a.txt"} and found["total"] == 2 and dict(found["views"])["duplicates"] == 2
    assert summary["same"] == 2 and counts["duplicates"] == 2
    assert sum("group by hash" in s for s in statements) == 3, statements  # 목록·폴더 설명·사이드바 숫자에 한 번씩


# B6 ─ 새로 지문을 만든 동영상만 다른 모든 동영상과 견준다. 기준(유사도)을 바꿨을 때만 모두 다시 맞춘다.
def fake_videos(db, count, seed=7):
    rng = random.Random(seed)
    bases = [[rng.getrandbits(64) for _ in range(12)] for _ in range(3)]  # 세 가지 내용
    for i in range(count):
        frames = [h ^ (1 << rng.randrange(64)) for h in rng.choice(bases)]  # 장면마다 1비트만 다르다
        db.execute("insert into media(id, kind, duration, frames, made, interval) values(?, 'video', ?, ?, 0, 30)",
                   (1000 + i, 100 + rng.random() * 30, media.encode(frames)))


def test_new_video_is_compared_only_with_the_others(home, monkeypatch):
    db = index.connect()
    try:
        fake_videos(db, 60)
        media.run(db)  # 처음: 모두 맞춘다
        everything = set(db.execute("select a, b, score from media_pair").fetchall())
        assert everything and db.execute("select value from meta where name = 'video_similar'").fetchone()[0] == 70
        new = 1000 + 59
        db.execute("delete from media_pair where a = ? or b = ?", (new, new))  # 지문을 새로 만든 것처럼(fingerprint가 짝을 지운다)
        db.execute("insert into media_todo(id) values(?)", (new,))
        calls = []
        real = media.similarity
        with monkeypatch.context() as patch:
            patch.setattr(media, "similarity", lambda a, b: calls.append(1) or real(a, b))
            media.run(db)
        assert set(db.execute("select a, b, score from media_pair").fetchall()) == everything  # 다 다시 맞춘 것과 같다
        assert 0 < len(calls) < 60 and db.execute("select count(*) from media_todo").fetchone()[0] == 0
        settings.save({"video_similar": 60})
        calls.clear()
        with monkeypatch.context() as patch:
            patch.setattr(media, "similarity", lambda a, b: calls.append(1) or real(a, b))
            media.run(db)
        assert len(calls) > 60  # 기준을 바꾸면 모두 다시
    finally:
        db.close()


def test_video_fingerprint_keeps_the_match_marker(home, monkeypatch):
    settings.save({"index_steps": [{"name": name, "on": False} for name in settings.STEPS]})
    write(home / "영상.mp4", "가짜 영상")
    build(home)
    db = index.connect()
    try:
        db.execute("insert or replace into meta(name, value) values('video_similar', 70)")
        row = dict(db.execute("select id, path, ext, stamp from files where name = '영상.mp4'").fetchone())
        with monkeypatch.context() as patch:
            patch.setattr(media, "video_length", lambda path: 30.0)
            patch.setattr(media, "frame_hashes", lambda path, seconds: [1, 2, 3, 4, 5])
            media.fingerprint(db, row, 30)
        assert db.execute("select value from meta where name = 'video_similar'").fetchone()[0] == 70
        assert db.execute("select count(*) from media_todo where id = ?", (row["id"],)).fetchone()[0] == 1
        media.match_images(db)  # 사진 짝 맞추기는 동영상 할 일을 지우지 않는다
        assert db.execute("select count(*) from media_todo where id = ?", (row["id"],)).fetchone()[0] == 1
    finally:
        db.close()


# B7 ─ 게임 묶음 폴더 이름(steam, blizzard 등)은 드라이브 바로 아래와 그 한 단계 아래에서만 게임 폴더로 본다.
def test_game_folder_names_deep_inside_are_ordinary_folders(home, client, monkeypatch):
    write(home / "steam" / "메모.txt")
    write(home / "Blizzard" / "내 사진" / "a.txt")
    build(home)
    assert names(client.post("/api/files", json={}).json["rows"]) == {"메모.txt", "a.txt"}
    assert games.found_games() == []
    with monkeypatch.context() as patch:
        patch.setattr(blacklist, "drive_roots", lambda roots: [str(home)])  # home을 드라이브처럼
        build(home)
        assert names(client.post("/api/files", json={}).json["rows"]) == set()
        assert [os.path.basename(g["path"]) for g in games.found_games()] == ["내 사진"]


# B8 ─ docx·xlsx·pptx도 MAX_CHARS까지만 읽고 멈춘다(PDF처럼).
def test_office_readers_stop_at_max_chars(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from pptx import Presentation
    word = docx(tmp_path / "긴.docx", [f"{n}번째 줄의 글자입니다" for n in range(300)])
    book = Workbook()
    for n in range(300):
        book.active.append([f"{n}번째 칸의 글자입니다"])
    book.save(tmp_path / "긴.xlsx")
    deck = Presentation()
    for n in range(100):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = f"{n}번째 장의 글자입니다"
    deck.save(tmp_path / "긴.pptx")
    with monkeypatch.context() as patch:
        patch.setattr(extract, "MAX_CHARS", 100)
        for read, path in ((extract.docx, word), (extract.xlsx, str(tmp_path / "긴.xlsx")), (extract.pptx, str(tmp_path / "긴.pptx"))):
            assert 100 <= len(read(path)) < 200, read


# B9 ─ 포트가 범위 밖인 주소는 '' (ValueError로 상세가 400이 되지 않게).
def test_bad_port_in_zone_url_is_masked(tmp_path):
    assert provenance.masked("https://example.com:99999/a.zip") == ""
    assert provenance.masked("https://user:pw@example.com:8080/a.zip?t=1") == "https://example.com:8080/a.zip?…"
    path = write(tmp_path / "받은.zip", "x")
    with open(path + ":Zone.Identifier", "w", encoding="utf-8") as stream:
        stream.write("[ZoneTransfer]\nZoneId=3\nHostUrl=https://example.com:99999/a.zip\n")
    assert provenance.info(path)["source"] == ""


# B10·R2(c)(d) ─ 끝에 점이 붙은 이름(\\?\로만 만들 수 있다)은 옆의 이름과 key가 같아 색인·정리·캡처에서 뺀다.
@pytest.fixture
def dotted(home):
    write(home / "abc", "옆의 다른 파일")
    odd = str(home / "abc.")
    with open("\\\\?\\" + odd, "w", encoding="utf-8") as stream:
        stream.write("끝에 점")
    yield odd
    os.remove("\\\\?\\" + odd)


def test_trailing_dot_name_is_not_indexed(home, dotted):
    assert [e.name for _, entries in index.walk(str(home)) for e in entries] == ["abc"]
    build(home)
    build(home)
    db = index.connect()
    try:
        assert [r[0] for r in db.execute("select path from files")] == [str(home / "abc")]
    finally:
        db.close()


def test_trailing_dot_name_is_not_organized_or_captured(home, dotted, tmp_path):
    dest = tmp_path / "놓을 곳"
    dest.mkdir()
    assert [os.path.basename(m["from"]) for m in organize.plan(str(home), str(dest))] == ["abc"]
    assert captures.files_in(str(home)) == [str(home / "abc")]


# B11 ─ 이름 검색은 파일 이름과 그 파일이 바로 든 폴더 이름에서만, 대소문자는 key와 같은 규칙으로 무시한다.
def test_name_search_matches_file_and_its_own_folder_only(home):
    write(home / "Desktop" / "sub" / "a.txt")
    write(home / "Ärger.txt")
    write(home / "Ölordner" / "b.txt")
    build(home)
    find = lambda name: names(search.listing({"name": name})["rows"])  # noqa: E731
    assert find("desktop") == set()  # 위 폴더 이름(사용자 이름·바탕 화면 등)으로는 찾지 않는다
    assert find(os.path.basename(os.path.dirname(str(home))).upper()) == set()
    assert find("SUB") == {"a.txt"} and find("ä") == {"Ärger.txt"} and find("ÖL") == {"b.txt"}


# B12 ─ 단계 목록에 사전이 아닌 것이 섞이면 400(ValueError).
def test_index_steps_with_non_dict_item_is_refused(home, client):
    steps = [{"name": name, "on": True} for name in settings.STEPS]
    with pytest.raises(ValueError, match="색인 단계 목록"):
        settings.check({"index_steps": steps + [1]})
    assert client.post("/api/settings", json={"index_steps": [1] + steps[1:]}).status_code == 400


# R1 ─ 크기 조건: 걸리는 파일이 적으면 크기 색인을 타고, 많으면 훑는다.
def test_narrow_size_condition_uses_size_index(home):
    for n in range(30):
        write(home / f"{n}.txt", "x")
    write(home / "큰.bin", "x" * (2 * 1024 * 1024))
    build(home)
    for query, uses_index in (({"size_min": "1"}, True), ({"size_min": "0.0000001"}, False), ({"size_min": "0", "size_max": "100"}, False)):
        parts, values, _ = search.conditions(query, settings.load())
        plan = plan_of("select f.path from files f" + search.where(parts), values)
        assert ("files_size" in plan) == uses_index, (query, plan)
    assert names(search.listing({"size_min": "1"})["rows"]) == {"큰.bin"}


# R2(a) ─ 늘 열 수 없는 사진(권한 없음)의 인식은 빈 결과로 기록한다(다시 시도하지 않게).
def test_unreadable_picture_is_recognized_as_empty(home, unlock_after):
    path = picture(home / "못읽음.png")
    build(home)
    sid = lock.user_sid()
    db = index.connect()
    try:
        row = dict(db.execute("select id, path, ext, stamp from files where name = '못읽음.png'").fetchone())
        subprocess.run(["icacls", path, "/deny", f"*{sid}:(R)"], check=True, capture_output=True)
        assert recognize.recognize_one(db, types.SimpleNamespace(np=None, faces=None, custom=None), row, {}, None)
        assert db.execute("select count(*) from recognized where id = ?", (row["id"],)).fetchone()[0] == 1
    finally:
        subprocess.run(["icacls", path, "/remove:d", f"*{sid}"], capture_output=True)
        db.close()


# R2(b) ─ 잠긴 Office 문서의 문서 비교 지문은 만들지 않고 다음에 다시 한다.
def test_locked_document_fingerprint_is_tried_later(home):
    path = docx(home / "잠긴.docx")
    build(home)
    db = index.connect()
    try:
        db.execute("update files set dochash = null, dochashed = null")
        handle = share_none(path)
        try:
            index.documents(db)
        finally:
            ctypes.windll.kernel32.CloseHandle(handle)
        assert db.execute("select dochashed from files where name = '잠긴.docx'").fetchone()[0] is None
        index.documents(db)
        assert db.execute("select dochashed from files where name = '잠긴.docx'").fetchone()[0] is not None
    finally:
        db.close()


# R5 ─ 사진을 읽다 늘 나는 오류(배드 섹터 등)는 빈 지문으로 기록하고, 잠시 막힌 것만 다시 한다.
def test_permanent_read_error_is_not_retried(tmp_path, monkeypatch):
    from PIL import Image
    path = picture(tmp_path / "배드섹터.png")
    for error, again in ((OSError(0, "데이터 오류(CRC)", None, 23), False), (OSError(0, "잠김", None, 33), True),
                         (PermissionError(13, "Permission denied"), True)):  # 마지막: C 런타임을 거쳐 번호가 빈 오류인데 파일은 읽힌다(잠깐 막혔다)
        def fail(*args, **kwargs):
            raise error

        with monkeypatch.context() as patch:
            patch.setattr(Image, "open", fail)
            if again:
                with pytest.raises(OSError):
                    media.image_hash(path)
            else:
                assert media.image_hash(path) is None
    assert media.image_hash(path) is not None
