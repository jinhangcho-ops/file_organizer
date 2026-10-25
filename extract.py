"""파일 종류별로 안의 글자를 꺼낸다(내용 검색용). 읽지 못하는 파일(암호, 깨짐)은 빈 글자."""
from itertools import chain
from pathlib import Path

MAX_SIZE = 50 * 1024 * 1024  # 이보다 큰 파일은 내용을 읽지 않는다(이름·크기 등으로만 찾는다).
MAX_CHARS = 2_000_000  # 파일 하나에서 저장할 최대 글자 수
PLAIN = {".txt", ".md", ".csv", ".tsv", ".log", ".json", ".xml", ".html", ".htm", ".ini", ".yaml", ".yml", ".py", ".sql"}


def plain(path):
    data = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "cp949"):  # 한글 Windows의 메모장 옛 파일은 cp949
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", "replace")


def gather(lines):
    """줄들을 MAX_CHARS가 찰 때까지만 모은다(큰 문서도 저장할 만큼만 읽고 멈춘다. 22MB xlsx는 끝까지 읽는 데 2분이 넘었다)."""
    kept, count = [], 0
    for line in lines:
        kept.append(line)
        count += len(line) + 1
        if count >= MAX_CHARS:
            break
    return "\n".join(kept)


def pdf(path):
    """pypdfium2(PDFium)로 읽는다. pypdf는 어떤 PDF에서 한 파일에 1분 넘게 걸렸다(PDFium은 1초)."""
    import pypdfium2

    document = pypdfium2.PdfDocument(path)
    try:
        lines, count = [], 0
        for page in document:
            textpage = page.get_textpage()
            lines.append(textpage.get_text_range())
            textpage.close()
            page.close()
            count += len(lines[-1])
            if count >= MAX_CHARS:
                break
        return "\n".join(lines)
    finally:
        document.close()


def docx(path):
    from docx import Document

    document = Document(path)
    tables = (" ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    return gather(chain((paragraph.text for paragraph in document.paragraphs), tables))


def xlsx(path):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return gather(" ".join(str(value) for value in row if value is not None)
                      for sheet in book.worksheets for row in sheet.iter_rows(values_only=True))
    finally:
        book.close()


def pptx(path):
    from pptx import Presentation

    return gather(shape.text_frame.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)


READERS = {".pdf": pdf, ".docx": docx, ".xlsx": xlsx, ".xlsm": xlsx, ".pptx": pptx} | {ext: plain for ext in PLAIN}


def readable(ext, size):
    return ext in READERS and size <= MAX_SIZE


def text(path, ext):
    """파일의 글자. 파일을 열거나 읽지 못하면(잠시 잠김, 클라우드 전용 등) None이라 색인이 다음에 다시 읽는다.
    내용이 깨진 파일과 늘 열 수 없는 파일(권한 없음)은 빈 글자(다시 읽어도 같으므로 읽은 것으로 둔다).
    먼저 한 번 열어 본다: PDFium·python-docx 등은 열지 못한 것을 제 오류(PdfiumError 등)로 바꿔 깨진 파일과 구별할 수 없다."""
    import media  # media가 index를 거쳐 이 파일을 불러오므로 여기서

    try:
        if not media.openable(path):
            return ""
        return READERS[ext](path)[:MAX_CHARS]
    except OSError:
        return None
    except Exception:  # 라이브러리마다 깨진 파일의 오류가 제각각이다.
        return ""
